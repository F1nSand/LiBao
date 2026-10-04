"""Build an editable LiBao thesis draft from the evidence-bound Markdown.

The existing reference DOCX is copied as a style carrier and never modified.
The content is read from docs/thesis/libao-thesis-draft.md so the Markdown and
DOCX remain easy to compare and revise.
"""

from __future__ import annotations

import os
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
REFERENCE = Path(
    os.environ.get(
        "LIBAO_THESIS_REFERENCE",
        r"C:\Users\Admin1\Desktop\开题报告-基于LLMAgent的二手车智能分析系统.docx",
    )
)
MARKDOWN = ROOT / "docs" / "thesis" / "libao-thesis-draft.md"
OUTPUT = ROOT / "docs" / "thesis" / "LiBao本科毕业论文初稿.docx"

sys.path.insert(0, str(ROOT / ".artifacts"))
import build_libao_proposal as style  # noqa: E402  reuse the inspected template contract


def clean_inline(text: str) -> str:
    text = re.sub(r"!\[([^]]*)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"\[([^]]+)\]\((https?://[^)]+)\)", r"\1（\2）", text)
    text = text.replace("**", "").replace("`", "")
    return text.strip()


def set_page_geometry(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.25)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    styles["Normal"].font.name = "宋体"
    styles["Normal"].font.size = Pt(12)
    styles["Normal"].paragraph_format.line_spacing = 1.275
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.first_line_indent = Inches(0.25)
    for name, size in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)):
        styles[name].font.name = "宋体"
        styles[name].font.size = Pt(size)
        styles[name].font.bold = True
        styles[name].paragraph_format.page_break_before = False


def add_cover(doc: Document) -> None:
    first = doc.sections[0]
    style.configure_footer(first, False)
    p = style.add_paragraph(
        doc,
        "本科毕业设计（论文）",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_indent=False,
        size=12,
        space_after=0,
    )
    p.paragraph_format.space_before = Pt(20)
    style.add_spacer(doc, 62)
    p = doc.add_paragraph(style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(28)
    r = p.add_run("LiBao：本地优先通用 Agent Runtime 与 Web 工作台的设计与实现")
    style.set_east_asia_font(r, size=18, bold=True)
    p = style.add_paragraph(
        doc,
        "——面向大语言模型智能体的工程化运行时研究",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_indent=False,
        size=12,
        space_after=0,
    )
    p.paragraph_format.space_after = Pt(32)
    rows = [
        ["论文题目", "LiBao：本地优先通用 Agent Runtime 与 Web 工作台的设计与实现"],
        ["学生姓名", "________________"],
        ["学    号", "________________"],
        ["院    系", "________________"],
        ["专    业", "________________"],
        ["指导教师", "________________"],
    ]
    style.add_table(doc, rows, [1700, 6940], header=False, font_size=12, cover=True)
    style.add_spacer(doc, 55)
    style.add_paragraph(
        doc,
        "20____年____月____日",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_indent=False,
        size=12,
        space_after=0,
    )


def add_body_section(doc: Document) -> None:
    second = doc.add_section(style.WD_SECTION.NEW_PAGE)
    second.page_width = Inches(8.27)
    second.page_height = Inches(11.69)
    second.left_margin = Inches(1.25)
    second.right_margin = Inches(1.25)
    second.top_margin = Inches(1)
    second.bottom_margin = Inches(1)
    style.configure_footer(second, True)
    style.set_page_number_start(second, 1)


def add_markdown_table(doc: Document, lines: list[str]) -> None:
    rows: list[list[str]] = []
    for line in lines:
        cells = [clean_inline(value.strip()) for value in line.strip().strip("|").split("|")]
        if cells and not all(set(cell) <= {"-", ":", " "} for cell in cells):
            rows.append(cells)
    if not rows:
        return
    widths = {
        2: [2500, 6140],
        3: [1500, 3000, 4140],
        4: [950, 2300, 2800, 2590],
    }.get(len(rows[0]), [8640 // len(rows[0])] * len(rows[0]))
    style.add_table(doc, rows, widths, header=True, font_size=10.5)
    style.add_spacer(doc, 4)


def add_code_block(doc: Document, lines: list[str]) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run("\n".join(lines))
    style.set_east_asia_font(r, name="Courier New", size=10.5)


def add_special_heading(doc: Document, text: str) -> None:
    p = style.add_paragraph(
        doc,
        clean_inline(text),
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_indent=False,
        size=14,
        space_after=8,
    )
    p.runs[0].bold = True


def add_markdown(doc: Document) -> None:
    lines = MARKDOWN.read_text(encoding="utf-8").splitlines()
    i = 0
    in_code = False
    code_lines: list[str] = []
    table_lines: list[str] = []
    chapter_started = False
    references_started = False
    content_started = False
    while i < len(lines):
        raw = lines[i]
        stripped = raw.strip()
        if not content_started:
            if stripped == "## 摘要":
                content_started = True
            else:
                i += 1
                continue
        if in_code:
            if stripped == "```":
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                code_lines.append(raw)
            i += 1
            continue
        if stripped.startswith("```"):
            if table_lines:
                add_markdown_table(doc, table_lines)
                table_lines = []
            in_code = True
            i += 1
            continue
        if stripped.startswith("|"):
            table_lines.append(raw)
            i += 1
            if i == len(lines) or not lines[i].strip().startswith("|"):
                add_markdown_table(doc, table_lines)
                table_lines = []
            continue
        if table_lines:
            add_markdown_table(doc, table_lines)
            table_lines = []
        if not stripped or stripped == "---":
            i += 1
            continue
        image_match = re.match(r"!\[([^]]*)\]\(([^)]+)\)", stripped)
        if image_match:
            image_path = (MARKDOWN.parent / image_match.group(2)).resolve()
            if image_path.is_file():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Inches(0)
                p.add_run().add_picture(str(image_path), width=Inches(6.0))
                caption = style.add_paragraph(
                    doc,
                    clean_inline(image_match.group(1)),
                    align=WD_ALIGN_PARAGRAPH.CENTER,
                    first_indent=False,
                    size=10.5,
                    space_after=8,
                )
                caption.runs[0].italic = True
            i += 1
            continue
        if stripped.startswith("# "):
            text = clean_inline(stripped[2:])
            if text.startswith("LiBao：") or text.startswith("本科毕业设计"):
                i += 1
                continue
            if text.startswith("第 "):
                if not chapter_started:
                    chapter_started = True
                heading = style.add_heading(doc, text, 1)
                if chapter_started:
                    heading.paragraph_format.page_break_before = True
            elif text == "参考文献":
                references_started = True
                style.add_heading(doc, text, 1)
            i += 1
            continue
        if stripped.startswith("## "):
            text = clean_inline(stripped[3:])
            if text in {"本科毕业设计（论文）初稿", "目录"}:
                if text == "目录":
                    add_special_heading(doc, text)
                i += 1
                continue
            if text in {"摘要", "Abstract"}:
                add_special_heading(doc, text)
            elif re.match(r"^\d+\.\d+\s", text):
                style.add_heading(doc, text, 2)
            else:
                style.add_heading(doc, text, 1)
            i += 1
            continue
        if stripped.startswith("### "):
            style.add_heading(doc, clean_inline(stripped[4:]), 2)
            i += 1
            continue
        if stripped.startswith("> "):
            p = style.add_paragraph(doc, clean_inline(stripped[2:]), first_indent=False, size=10.5, space_after=6)
            for run in p.runs:
                run.italic = True
            i += 1
            continue
        if re.match(r"^\[\d+\]", stripped) and references_started:
            style.add_reference(doc, clean_inline(stripped))
            i += 1
            continue
        list_match = re.match(r"^(?:[-*]|\d+\.)\s+(.*)$", stripped)
        if list_match:
            numbered = bool(re.match(r"^\d+\.", stripped))
            style.add_list(doc, clean_inline(list_match.group(1)), numbered=numbered)
            i += 1
            continue
        if stripped.startswith("!["):
            i += 1
            continue
        style.add_paragraph(doc, clean_inline(stripped), first_indent=True, size=12, space_after=6)
        i += 1


def build() -> Path:
    if not REFERENCE.is_file():
        raise FileNotFoundError(
            f"reference DOCX not found: {REFERENCE}. "
            "Set LIBAO_THESIS_REFERENCE to the template path."
        )
    doc = Document(str(REFERENCE))
    style.clear_body(doc)
    set_page_geometry(doc)
    configure_styles(doc)
    add_cover(doc)
    add_body_section(doc)
    add_markdown(doc)
    doc.save(str(OUTPUT))
    return OUTPUT


if __name__ == "__main__":
    print(build())
